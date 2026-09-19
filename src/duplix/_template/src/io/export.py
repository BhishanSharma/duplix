"""Build a downloadable .xlsx from the current state.

The app no longer keeps a workbook — this module exists only so the
assigner can click "Download .xlsx" and walk away with a file to email
or print. It is a pure read of ``AppState``: nothing here is ever read
back in.

One sheet per output, plus the colour cues that carried meaning in the
old console workbook (international DEP, pre-plan / relief handovers,
warnings, unallocated rows).
"""

from __future__ import annotations

import io
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..schemas import AllocationRow, AllocationSheet, OpsClass
from ..state import AppState, OVERRIDE_HEADERS

# Colour cues, carried over from the console workbook.
HEADER_FILL_HEX = "FF1F3864"
INTL_FILL_HEX = "FFFFD6E7"        # soft pink — international DEP
PLANNED_BY_FILL_HEX = "FFDEEAF6"  # blue — pre-planned by someone else
RELIEVED_BY_FILL_HEX = "FFE2EFDA" # green — handed over mid-flight
WARNING_FILL_HEX = "FFFFF2CC"     # amber
UNALLOCATED_FILL_HEX = "FFFCE4E4" # soft red

ALLOCATION_HEADERS = (
    "date", "flt", "dep", "arr", "std", "pax",
    "staff", "planned_by", "relieved_by", "warning",
)
CLEANED_HEADERS = (
    "date", "flt", "type", "ac", "dep", "arr", "std", "pax", "ops_class",
)
PAIR_MAP_HEADERS = (
    "boundary", "pair_role", "preplan_count",
    "prev_employee_id", "prev_name", "prev_shift",
    "next_employee_id", "next_name", "next_shift",
)
UNALLOCATED_HEADERS = (
    "date", "flt", "dep", "arr", "std", "pax",
    "ops_class", "is_international", "reason",
)
WORKLOAD_HEADERS = (
    "employee_id", "name", "shift", "role",
    "target_preferred", "target_acceptable_max", "hard_cap", "actual",
    "deviation_below_preferred", "deviation_above_preferred", "violations",
)
WARNING_HEADERS = ("severity", "code", "name", "date", "message")
AVAILABILITY_HEADERS = (
    "employee_id", "name", "role", "date", "status", "raw_status",
    "assignable", "current_shift", "license",
)

_ALLOC_SHEET_NAME = {
    AllocationSheet.DAY_OPS: "Allocations_DayOps",
    AllocationSheet.NIGHT_OPS: "Allocations_NightOps",
    AllocationSheet.P2F: "Allocations_P2F",
    AllocationSheet.NORSE: "Allocations_NORSE",
}
_CLEANED_SHEET_NAME = {
    OpsClass.DAY: "Flights_DayOps",
    OpsClass.NIGHT: "Flights_NightOps",
    OpsClass.P2F: "Flights_P2F",
    OpsClass.NORSE: "Flights_NORSE",
    OpsClass.GULF: "Flights_Gulf",
    OpsClass.TEST: "Flights_Test",
    OpsClass.FERRY: "Flights_Ferry",
    OpsClass.CHARTER: "Flights_Charter",
}

# Column indices (1-based) used for the per-cell fills.
_COL_DEP = ALLOCATION_HEADERS.index("dep") + 1
_COL_ARR = ALLOCATION_HEADERS.index("arr") + 1
_COL_PLANNED_BY = ALLOCATION_HEADERS.index("planned_by") + 1
_COL_RELIEVED_BY = ALLOCATION_HEADERS.index("relieved_by") + 1
_COL_WARNING = ALLOCATION_HEADERS.index("warning") + 1


def _sheet(wb: Workbook, title: str, headers: Iterable[str]) -> Worksheet:
    """Create a sheet with a styled, frozen header row."""
    ws = wb.create_sheet(title[:31])
    cols = list(headers)
    ws.append(cols)
    fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
    for col in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(
            12, min(48, len(str(cols[col - 1])) + 6),
        )
    ws.freeze_panes = "A2"
    return ws


def _fill(ws: Worksheet, row: int, col: int, hex_color: str) -> None:
    ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=hex_color)


def _write_summary(wb: Workbook, state: AppState) -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    rows = [
        ("Allocation date", state.run_date.isoformat() if state.run_date else "—"),
        ("Last run", state.last_run or "—"),
        ("Mode", state.mode or "—"),
        ("Solver status", state.solver_status or "—"),
        ("Duration (s)", round(state.duration_s, 1) if state.duration_s else "—"),
        ("", ""),
        ("Flights allocated", len(state.allocations)),
        ("Flights unallocated", len(state.unallocated)),
        ("Staff in workload summary", len(state.workload)),
        ("Pairs", len(state.pairs)),
        ("Warnings", len(state.warnings)),
        ("", ""),
        ("Input files", ""),
    ]
    for label, value in rows:
        ws.append([label, value])
        if label and not isinstance(value, (int, float)):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for up in state.inputs_as_json():
        ws.append([
            f"  {up['label']}",
            up["filename"] or "(not uploaded)",
        ])
    if state.plan_text:
        ws.append(["", ""])
        ws.append(["Plan summary", ""])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for line in state.plan_text.splitlines():
            if line.strip():
                ws.append(["", line])


def _write_allocations(wb: Workbook, state: AppState) -> None:
    by_sheet: dict[AllocationSheet, list[AllocationRow]] = {
        target: [] for target in _ALLOC_SHEET_NAME
    }
    for r in state.allocations:
        by_sheet.setdefault(r.sheet_target, []).append(r)
    for target, name in _ALLOC_SHEET_NAME.items():
        ws = _sheet(wb, name, ALLOCATION_HEADERS)
        for r in by_sheet.get(target, []):
            ws.append([
                r.date.isoformat(),
                r.flt,
                r.dep,
                r.arr,
                r.std.isoformat(timespec="minutes"),
                r.pax,
                r.staff_name,
                r.planned_by_name or "",
                r.relieved_by_name or "",
                r.warning or "",
            ])
            row_idx = ws.max_row
            if r.planned_by_employee_id:
                _fill(ws, row_idx, _COL_PLANNED_BY, PLANNED_BY_FILL_HEX)
            if r.relieved_by_employee_id:
                _fill(ws, row_idx, _COL_RELIEVED_BY, RELIEVED_BY_FILL_HEX)
            if r.warning:
                _fill(ws, row_idx, _COL_WARNING, WARNING_FILL_HEX)
            if r.is_international:
                _fill(ws, row_idx, _COL_DEP, INTL_FILL_HEX)
                _fill(ws, row_idx, _COL_ARR, INTL_FILL_HEX)


def _write_cleaned(wb: Workbook, state: AppState) -> None:
    for ops, name in _CLEANED_SHEET_NAME.items():
        rows = state.cleaned.get(ops, [])
        if not rows:
            continue
        ws = _sheet(wb, name, CLEANED_HEADERS)
        for r in rows:
            ws.append([
                r.date.isoformat(), r.flt, r.type, r.ac, r.dep, r.arr,
                r.std.isoformat(timespec="minutes"), r.load, r.ops_class.value,
            ])
            if r.is_international:
                _fill(ws, ws.max_row, CLEANED_HEADERS.index("dep") + 1,
                      INTL_FILL_HEX)


def build_workbook(state: AppState) -> bytes:
    """Render the whole state as a single .xlsx and return its bytes."""
    wb = Workbook()
    # Workbook() ships with one default sheet we don't want.
    wb.remove(wb.active)

    with state.lock:
        _write_summary(wb, state)
        _write_allocations(wb, state)

        ws = _sheet(wb, "Unallocated", UNALLOCATED_HEADERS)
        for row in state.unallocated:
            ws.append([
                row.get("date", ""), row.get("flt", ""), row.get("dep", ""),
                row.get("arr", ""), row.get("std", ""), row.get("pax", 0),
                row.get("ops_class", ""),
                "Y" if row.get("is_international") else "N",
                row.get("reason", ""),
            ])
            for col in range(1, len(UNALLOCATED_HEADERS) + 1):
                _fill(ws, ws.max_row, col, UNALLOCATED_FILL_HEX)

        ws = _sheet(wb, "Workload", WORKLOAD_HEADERS)
        for r in state.workload:
            ws.append([
                r.employee_id, r.name, r.shift or "", r.role.value,
                r.target_preferred, r.target_acceptable_max, r.hard_cap,
                r.actual, r.deviation_below_preferred,
                r.deviation_above_preferred, "; ".join(r.violations),
            ])

        ws = _sheet(wb, "Pair_Map", PAIR_MAP_HEADERS)
        for p in state.pairs:
            ws.append([
                p.boundary.value, p.pair_role.value, p.preplan_count,
                p.prev_employee_id or "", p.prev_name or "", p.prev_shift or "",
                p.next_employee_id or "", p.next_name or "", p.next_shift or "",
            ])

        ws = _sheet(wb, "Warnings", WARNING_HEADERS)
        for w in state.warnings:
            ws.append([
                w.severity.value, w.code, w.name or "",
                w.date.isoformat() if w.date else "", w.message,
            ])
            if w.severity.value in ("WARN", "ERROR"):
                for col in range(1, len(WARNING_HEADERS) + 1):
                    _fill(ws, ws.max_row, col, WARNING_FILL_HEX)

        _write_cleaned(wb, state)

        ws = _sheet(wb, "Roster", AVAILABILITY_HEADERS)
        for av in state.availability:
            ws.append([
                av.employee_id, av.name, av.role.value, av.date.isoformat(),
                av.status.value, av.raw_status, av.assignable,
                av.current_shift or "", av.license or "",
            ])

        if state.staffing:
            headers = list(state.staffing[0].keys())
            ws = _sheet(wb, "Staffing_Recommendation", headers)
            for row in state.staffing:
                ws.append([row.get(h, "") for h in headers])

        if state.overrides:
            ws = _sheet(wb, "Overrides", OVERRIDE_HEADERS)
            for row in state.overrides:
                ws.append([row.get(h, "") for h in OVERRIDE_HEADERS])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def filename_for(state: AppState) -> str:
    """A download name that carries the allocation date."""
    d = state.run_date.isoformat() if state.run_date else "undated"
    return f"flight_allocation_{d}.xlsx"
