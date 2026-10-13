"""Turn any accepted upload into an openpyxl workbook.

The SV portal and the rostering tools do not always hand over a real
.xlsx: an "Excel export" is often a tab-separated text file or an HTML
table saved with an .xls name, and some people still save the old binary
.xls. Every reader in the engine works on an openpyxl ``Workbook``, so
this module is the single place that recognises the file by its CONTENT
(not its name) and converts it:

    .xlsx / .xlsm             zip            openpyxl, as is
    .xls (Excel 97-2003)      OLE2 binary    xlrd
    "xls" that is HTML        <table> text   html.parser
    .csv / .tsv / .txt / "xls" that is text  csv (delimiter sniffed)

``to_xlsx_bytes`` is applied at upload time, so what the server keeps
(and the roster store writes to disk) is always a real .xlsx.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from html.parser import HTMLParser

import openpyxl
from openpyxl import Workbook

ACCEPTED_TYPES = ".xlsx, .xlsm, .xls, .csv, .tsv, .txt"

_ZIP = b"PK"
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class UnreadableUpload(ValueError):
    """The bytes are not any format we can turn into a workbook."""


def sniff(data: bytes) -> str:
    """'xlsx' | 'xls' | 'html' | 'text' — decided by content."""
    if data.startswith(_ZIP):
        return "xlsx"
    if data.startswith(_OLE2):
        return "xls"
    head = data[:2048].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if head.startswith((b"<!doctype html", b"<html", b"<table", b"<?xml")) or b"<table" in head:
        return "html"
    return "text"


def to_workbook(data: bytes, *, read_only: bool = False) -> Workbook:
    """Open ``data`` as a workbook, whatever accepted format it is in.

    ``read_only`` opens an .xlsx lazily (rows are parsed only as they
    are iterated): much faster when the caller only peeks at it."""
    if not data:
        raise UnreadableUpload("the file is empty")
    kind = sniff(data)
    try:
        if kind == "xlsx":
            return openpyxl.load_workbook(
                io.BytesIO(data), data_only=True, read_only=read_only,
            )
        if kind == "xls":
            return _from_xls(data)
        if kind == "html":
            return _from_rows_by_sheet(_html_tables(_decode(data)))
        return _from_rows_by_sheet({"Sheet1": _text_rows(_decode(data))})
    except UnreadableUpload:
        raise
    except Exception as exc:  # noqa: BLE001 — becomes a user-facing message
        raise UnreadableUpload(
            f"could not read it as {kind} ({type(exc).__name__}: {exc})"
        ) from exc


def to_xlsx_bytes(data: bytes) -> bytes:
    """``data`` as real .xlsx bytes (returned unchanged if it already is)."""
    if sniff(data) == "xlsx":
        to_workbook(data, read_only=True).close()   # still validate it opens
        return data
    wb = to_workbook(data)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- converters ----------

def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-16" and "\x00" in text:
            continue
        return text
    raise UnreadableUpload("unknown text encoding")


def _text_rows(text: str) -> list[list[object]]:
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t,;|")
        delim = dialect.delimiter
    except csv.Error:
        # Pick the separator that splits the first line the most.
        first = sample.splitlines()[0] if sample.splitlines() else ""
        delim = max("\t,;|", key=first.count)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not any(any(c.strip() for c in r) for r in rows):
        raise UnreadableUpload("the text file has no data")
    return [[_text_cell(c) for c in r] for r in rows]


def _text_cell(v: str) -> object:
    """Text cells stay text, except plain whole numbers (IDs, pax)."""
    s = v.strip()
    if not s:
        return None
    if s.isdigit() and len(s) < 15 and not (len(s) > 1 and s.startswith("0")):
        return int(s)
    return s


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _html_tables(text: str) -> dict[str, list[list[object]]]:
    p = _TableParser()
    p.feed(text)
    tables = [t for t in p.tables if t]
    if not tables:
        raise UnreadableUpload("no table found in the HTML file")
    return {
        f"Sheet{i + 1}": [[_text_cell(c) for c in row] for row in t]
        for i, t in enumerate(sorted(tables, key=len, reverse=True))
    }


def _from_xls(data: bytes) -> Workbook:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover — listed in requirements
        raise UnreadableUpload(
            "old-style .xls needs the 'xlrd' package — re-save the file as "
            ".xlsx, or install it once with "
            r"'.venv\Scripts\pip install -r requirements.txt'"
        ) from exc
    book = xlrd.open_workbook(file_contents=data)
    sheets: dict[str, list[list[object]]] = {}
    for sh in book.sheets():
        rows: list[list[object]] = []
        for r in range(sh.nrows):
            row: list[object] = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate_as_datetime(v, book.datemode)
                    # Pure time-of-day cells come back on 1899-12-31.
                    v = dt.time() if dt.year < 1901 else dt
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                    v = int(v)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                elif isinstance(v, str):
                    v = v.strip() or None
                row.append(v)
            rows.append(row)
        sheets[sh.name] = rows
    return _from_rows_by_sheet(sheets)


def _from_rows_by_sheet(sheets: dict[str, list[list[object]]]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=_safe_sheet_name(name, wb.sheetnames))
        for row in rows:
            ws.append(list(row))
    if not wb.sheetnames:
        raise UnreadableUpload("no sheets found")
    wb.properties.modified = datetime.now()
    return wb


def _safe_sheet_name(name: str, taken: list[str]) -> str:
    clean = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
    out, n = clean, 2
    while out in taken:
        out = f"{clean[:28]}_{n}"
        n += 1
    return out
